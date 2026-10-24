# -*- coding: utf-8 -*-
import os
import sys

import pytest
import yaml

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

CFG_PATH = os.path.join(ROOT, "config", "employee_zhoubao.yaml")


@pytest.fixture(scope="module")
def cfg():
    with open(CFG_PATH, "r", encoding="utf-8") as f:
        c = yaml.safe_load(f)
    from app.tools import registry
    registry.set_policy(c["tools_allowed"], c["limits"]["max_tool_retry"])
    return c


@pytest.fixture(scope="module")
def data(tmp_path_factory):
    """生成示例销售明细（xlsx）与周会纪要（docx）到临时目录。"""
    import pandas as pd
    from docx import Document
    from scripts.make_sample_data import build_sales

    d = tmp_path_factory.mktemp("sample")
    rows = build_sales()
    sales_p = os.path.join(str(d), "销售明细.xlsx")
    pd.DataFrame(rows)[["日期", "门店", "品类", "销售额", "订单数"]].to_excel(
        sales_p, index=False)
    doc = Document()
    doc.add_heading("周会纪要", level=1)
    doc.add_paragraph("福田店本周三因冷萃原料到货延迟，当日出现部分商品断货，销售额出现明显下滑。")
    doc.add_paragraph("南山店在周末推出第二杯半价活动，南山店客流与销售额明显抬升。")
    notes_p = os.path.join(str(d), "周会纪要.docx")
    doc.save(notes_p)
    return {"sales": sales_p, "notes": notes_p}


@pytest.fixture
def ws(tmp_path):
    return str(tmp_path)
